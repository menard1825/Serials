from dataclasses import dataclass, field
from datetime import date
from io import BytesIO

import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from hardware_helpers import HardwareItem, find_duplicates, hardware_serial_count, parse_bulk_hardware, parse_pasted_values, safe_filename_part

SAFARI_BLUE = "004785"
SAFARI_LIGHT_BLUE = "EAF2F8"
WHITE = "FFFFFF"


@dataclass
class SoftwareItem:
    software: str
    license_type: str = ""
    values: list[str] = field(default_factory=list)


def clean_hardware(items):
    cleaned = []
    for item in items:
        serials = [s.strip() for s in item.serials if s.strip()]
        if item.product.strip() and serials:
            cleaned.append(HardwareItem(item.product.strip(), item.model.strip(), serials))
    return cleaned


def clean_software(items):
    cleaned = []
    for item in items:
        values = [value.strip() for value in item.values if value.strip()]
        if item.software.strip() and values:
            cleaned.append(SoftwareItem(item.software.strip(), item.license_type.strip(), values))
    return cleaned


def software_value_count(items):
    return sum(len(item.values) for item in items)


def _table_header(ws, row, labels, border):
    for col, label in enumerate(labels, 1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=SAFARI_BLUE)
        cell.border = border
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 22


def _data_row(ws, row, values, border, fill=None):
    for col, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)


def create_customer_report(client_name, order_number, client_po, hardware_items, software_items):
    wb = Workbook()
    ws = wb.active
    ws.title = "Serials & Licenses"
    ws.sheet_view.showGridLines = False
    thin = Side(style="thin", color="C9D1D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:C2")
    ws["A1"] = "SAFARI MICRO\nSerial Numbers & License Report"
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(vertical="center", wrap_text=True)
    for cells in ws["A1:C2"]:
        for cell in cells:
            cell.fill = PatternFill("solid", fgColor=SAFARI_BLUE)
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 28

    details = [
        ("Client", client_name or "—"),
        ("Safari Order #", order_number or "—"),
        ("Customer PO", client_po or "—"),
        ("Prepared", date.today().strftime("%B %d, %Y").replace(" 0", " ")),
    ]
    row = 4
    for label, value in details:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        row += 1

    hw_total = hardware_serial_count(hardware_items)
    sw_total = software_value_count(software_items)
    summary = []
    if hw_total:
        summary.append(f"{hw_total} hardware serial number{'s' if hw_total != 1 else ''}")
    if sw_total:
        summary.append(f"{sw_total} software license entr{'ies' if sw_total != 1 else 'y'}")
    ws.merge_cells("A9:C9")
    ws["A9"] = "Report includes " + " and ".join(summary)
    ws["A9"].fill = PatternFill("solid", fgColor=SAFARI_LIGHT_BLUE)
    ws["A9"].font = Font(italic=True)

    row = 11
    first_header = None
    if hardware_items:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        ws.cell(row=row, column=1, value="Hardware Serial Numbers").font = Font(size=14, bold=True, color=SAFARI_BLUE)
        row += 1
        first_header = row
        _table_header(ws, row, ["Product", "Model Number", "Serial Number"], border)
        row += 1
        for index, item in enumerate(hardware_items):
            fill = SAFARI_LIGHT_BLUE if index % 2 else None
            for serial in item.serials:
                _data_row(ws, row, [item.product, item.model, serial], border, fill)
                row += 1
        row += 2

    if software_items:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        ws.cell(row=row, column=1, value="Software License Information").font = Font(size=14, bold=True, color=SAFARI_BLUE)
        row += 1
        first_header = first_header or row
        _table_header(ws, row, ["Software", "License Type", "License / Activation Code"], border)
        row += 1
        for index, item in enumerate(software_items):
            fill = SAFARI_LIGHT_BLUE if index % 2 else None
            for value in item.values:
                _data_row(ws, row, [item.software, item.license_type, value], border, fill)
                row += 1
        row += 2

    ws.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=3)
    ws.cell(row=row, column=1, value="Thank you for your business. If you need assistance with this report, please contact your Safari Micro representative.")
    ws.cell(row=row, column=1).font = Font(size=10, italic=True, color="5F6B73")
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 42
    ws.freeze_panes = f"A{(first_header or 10) + 1}"
    ws.print_title_rows = f"1:{min(first_header or 10, 12)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def _init_state():
    st.session_state.setdefault("hardware_count", 1)
    st.session_state.setdefault("software_count", 1)
    st.session_state.setdefault("include_hardware", True)
    st.session_state.setdefault("include_software", False)


def _reset_report():
    keep = {"include_hardware": True, "include_software": False, "hardware_count": 1, "software_count": 1}
    for key in list(st.session_state.keys()):
        if key.startswith(("client_", "order_", "po_", "hw_", "sw_", "bulk_", "generated_")):
            st.session_state.pop(key, None)
    for key, value in keep.items():
        st.session_state[key] = value


def _render_hardware():
    st.subheader("Hardware serial numbers")
    method = st.radio("Hardware entry", ["Easy entry", "Bulk paste from Excel"], horizontal=True, label_visibility="collapsed", key="hw_entry_method")
    if method == "Bulk paste from Excel":
        st.caption("Copy rows from Excel in this order: Product | Model | Serial Number. Two-column Product | Serial pastes also work.")
        raw = st.text_area("Paste hardware rows", key="bulk_hardware", height=220, placeholder="Dell Latitude 7450\tLatitude 7450\tABC1234")
        items, ignored = parse_bulk_hardware(raw)
        count = hardware_serial_count(items)
        if raw.strip() and count:
            st.success(f"{count} serial number{'s' if count != 1 else ''} detected across {len(items)} product{'s' if len(items) != 1 else ''}.")
        if ignored:
            st.warning(f"{len(ignored)} row{'s were' if len(ignored) != 1 else ' was'} skipped. Check the pasted columns.")
        return items

    items = []
    for i in range(st.session_state.hardware_count):
        with st.container(border=True):
            st.markdown(f"**Hardware item {i + 1}**")
            c1, c2 = st.columns(2)
            with c1:
                product = st.text_input("Product", key=f"hw_product_{i}", placeholder="Dell Latitude 7450")
            with c2:
                model = st.text_input("Model number", key=f"hw_model_{i}", placeholder="Latitude 7450")
            raw = st.text_area("Serial numbers", key=f"hw_serials_{i}", height=120, placeholder="Paste one per line, from Excel, or comma-separated...")
            serials = parse_pasted_values(raw)
            duplicates = find_duplicates(serials)
            if serials and not duplicates:
                st.caption(f"✓ {len(serials)} serial number{'s' if len(serials) != 1 else ''} detected")
            if duplicates:
                st.warning("Duplicate serial number" + ("s" if len(duplicates) != 1 else "") + ": " + ", ".join(duplicates[:8]))
            items.append(HardwareItem(product, model, serials))

    add_col, remove_col, _ = st.columns([1.5, 1.2, 4])
    with add_col:
        if st.button("+ Add another product", use_container_width=True):
            st.session_state.hardware_count += 1
            st.rerun()
    with remove_col:
        if st.session_state.hardware_count > 1 and st.button("Remove last", use_container_width=True):
            st.session_state.hardware_count -= 1
            st.rerun()
    return items


def _render_software():
    st.subheader("Software license information")
    items = []
    for i in range(st.session_state.software_count):
        with st.container(border=True):
            st.markdown(f"**Software item {i + 1}**")
            c1, c2 = st.columns(2)
            with c1:
                software = st.text_input("Software", key=f"sw_name_{i}", placeholder="Microsoft Office")
            with c2:
                license_type = st.text_input("License type", key=f"sw_type_{i}", placeholder="Perpetual / Subscription / MAK")
            raw = st.text_area("License / activation codes", key=f"sw_values_{i}", height=120, placeholder="Paste one per line, from Excel, or comma-separated...")
            values = parse_pasted_values(raw)
            duplicates = find_duplicates(values)
            if values and not duplicates:
                st.caption(f"✓ {len(values)} entr{'ies' if len(values) != 1 else 'y'} detected")
            if duplicates:
                st.info("Repeated software value detected. This can be valid for some licensing programs, so confirm before sending.")
            items.append(SoftwareItem(software, license_type, values))

    add_col, remove_col, _ = st.columns([1.6, 1.4, 4])
    with add_col:
        if st.button("+ Add another software item", use_container_width=True):
            st.session_state.software_count += 1
            st.rerun()
    with remove_col:
        if st.session_state.software_count > 1 and st.button("Remove last software", use_container_width=True):
            st.session_state.software_count -= 1
            st.rerun()
    return items


def main():
    st.set_page_config(page_title="Safari Micro | Serials & Licenses", page_icon="📄", layout="wide")
    _init_state()
    st.image("https://safarimicro.com/wp-content/uploads/2022/01/SafariMicro-Color-with-Solid-Icon-Copy.png", width=230)
    st.title("Serial Numbers & License Report")
    st.caption("Create a clean, customer-ready Excel report without fighting with formatting.")

    st.subheader("1. Customer information")
    c1, c2, c3 = st.columns(3)
    with c1:
        client_name = st.text_input("Client name", key="client_name")
    with c2:
        order_number = st.text_input("Safari order #", key="order_number")
    with c3:
        client_po = st.text_input("Customer PO", key="po_number")

    st.subheader("2. What are you sending?")
    c1, c2, _ = st.columns([1.5, 1.5, 4])
    with c1:
        include_hardware = st.checkbox("Hardware serial numbers", key="include_hardware")
    with c2:
        include_software = st.checkbox("Software license information", key="include_software")

    hardware = _render_hardware() if include_hardware else []
    if include_hardware:
        st.divider()
    software = _render_software() if include_software else []

    hardware = clean_hardware(hardware)
    software = clean_software(software)
    hw_count = hardware_serial_count(hardware)
    sw_count = software_value_count(software)
    hw_duplicates = find_duplicates([serial for item in hardware for serial in item.serials])

    st.divider()
    st.subheader("3. Create the report")
    problems = []
    if not client_name.strip():
        problems.append("Add the client name.")
    if not order_number.strip():
        problems.append("Add the Safari order number.")
    if not include_hardware and not include_software:
        problems.append("Select hardware, software, or both.")
    if include_hardware and not hw_count:
        problems.append("Add at least one hardware serial number.")
    if include_software and not sw_count:
        problems.append("Add at least one software entry.")
    if hw_duplicates:
        problems.append("Remove duplicate hardware serial numbers: " + ", ".join(hw_duplicates[:8]))

    if problems:
        st.warning("Please finish the items below before creating the report:\n\n" + "\n".join(f"• {item}" for item in problems))
    else:
        parts = []
        if hw_count:
            parts.append(f"{hw_count} hardware serial number{'s' if hw_count != 1 else ''}")
        if sw_count:
            parts.append(f"{sw_count} software entr{'ies' if sw_count != 1 else 'y'}")
        st.success("Report looks good — " + " and ".join(parts) + ".")

    left, middle, _ = st.columns([1.8, 1, 5])
    with left:
        generate = st.button("Create Excel report", type="primary", use_container_width=True, disabled=bool(problems))
    with middle:
        st.button("Start over", on_click=_reset_report, use_container_width=True)

    if generate:
        st.session_state.generated_excel = create_customer_report(client_name, order_number, client_po, hardware, software)
        st.session_state.generated_filename = f"SafariMicro_{safe_filename_part(client_name, 'Client')}_{safe_filename_part(order_number, 'Order')}_Serials.xlsx"

    if st.session_state.get("generated_excel"):
        st.download_button("Download customer report", data=st.session_state.generated_excel, file_name=st.session_state.generated_filename, mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
        st.caption("The report is generated in memory for download; this app does not write the entered report data to a database or local file.")


if __name__ == "__main__":
    main()
